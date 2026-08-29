from __future__ import annotations

import hashlib
import html
import os
import re
import threading
import unicodedata
import uuid
from pathlib import Path
from typing import Any

import pymupdf as fitz

from .db import connect, project_id, utc_now
from .library_store import connect_library, initialize_library, resolve_library_root
from .skill_registry import discover_skills, get_skill
from .content_graph import project_content_graph


SUPPORTED_SUFFIXES = {
    ".pdf", ".md", ".txt", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv",
    ".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff",
    ".geojson", ".gpkg", ".kml", ".kmz", ".mbtiles",
}
CANDIDATE_SUFFIXES = SUPPORTED_SUFFIXES | {".doc", ".epub", ".caj"}
HISTORY_TERMS = (
    "历史", "史料", "档案", "地方志", "编年", "朝代", "帝国", "革命", "战争", "考察",
    "history", "historical", "archive", "chronicle", "century", "empire", "revolution", "war",
)
LIBRARY_SHELVES = {
    "primary_sources": "原始史料", "academic_articles": "学术论文",
    "monographs": "学术专著", "personal_manuscripts": "个人论文与稿件",
    "reading_notes": "读书笔记", "reference_works": "工具书与目录", "unclassified": "待分类",
}
GRAPH_WORK_SHELVES = frozenset({
    "primary_sources", "academic_articles", "monographs",
    "personal_manuscripts", "reference_works",
})
SUGGESTED_SHELVES = {
    "archival_source": "primary_sources",
    "article": "academic_articles",
    "thesis": "academic_articles",
    "monograph": "monographs",
    "personal_manuscript": "personal_manuscripts",
    "reading_note": "reading_notes",
    "reference_work": "reference_works",
}
SCAN_PAGE_SIZE = 50
SCAN_MAX_PAGE_SIZE = 50
SCAN_WRITE_BATCH_SIZE = 50


def _automatic_scan_exclusion(path: Path, source_root: Path) -> str:
    """Keep generated benchmarks and transient engineering files out of library intake."""
    try:
        relative = path.resolve().relative_to(source_root.resolve())
    except ValueError:
        relative = path
    parts = [part.casefold() for part in relative.parts]
    blocked_parts = {".git", ".codex-work", "node_modules", "__pycache__", "pytest_tmp", ".pytest_cache"}
    if any(part in blocked_parts or part.startswith("pytest-") for part in parts[:-1]):
        return "工程缓存或测试运行目录"
    stem = path.stem.casefold()
    generated_terms = ("histra-bench", "benchmark", "mcqtask", "final_three_questions")
    if any(term in stem for term in generated_terms) or re.search(r"(?:^|[_\s-])bench(?:$|[_\s-])", stem):
        return "Bench 或自动评测材料"
    if path.suffix.casefold() in {".md", ".txt"} and (
        stem in {"file", "readme", "agents", "skill", "changelog"}
        or stem.startswith(("cache_", "codex-clipboard-"))
    ):
        return "临时 Markdown、文本或工程说明"
    return ""


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _graph_node(connection: Any, node_type: str, label: str, normalized: str | None = None) -> str:
    normalized_label = (normalized or label).strip().casefold()
    node_id = "KGN_" + hashlib.sha256(f"{node_type}\0{normalized_label}".encode("utf-8")).hexdigest()[:24]
    connection.execute(
        """INSERT INTO knowledge_nodes(node_id, node_type, label, normalized_label, origin, created_at)
           VALUES (?, ?, ?, ?, 'bibliographic_metadata', ?)
           ON CONFLICT(node_type, normalized_label) DO UPDATE SET label = excluded.label""",
        (node_id, node_type, label.strip(), normalized_label, utc_now()),
    )
    return node_id


def _author_names(value: str) -> list[str]:
    cleaned = re.sub(r"\s*[（(](?:著|编|主编|译|校|dir\.?|ed\.?|editor)[^）)]*[）)]", "", str(value or ""), flags=re.IGNORECASE)
    parts = re.split(r"[；;,，、]|\s+(?:&|and|et|und|和)\s+", cleaned, flags=re.IGNORECASE)
    return list(dict.fromkeys(part.strip() for part in parts if _clean_author(part.strip())))


def _author_key(value: str) -> str:
    return re.sub(r"[\s.．·•'’`_-]+", "", unicodedata.normalize("NFKC", value).casefold())


def _author_identity(connection: Any, value: str) -> tuple[str, str]:
    key = _author_key(value)
    row = connection.execute(
        "SELECT canonical_name FROM author_aliases WHERE alias_normalized=?", (key,)
    ).fetchone()
    canonical = str(row["canonical_name"]) if row else value
    return canonical, _author_key(canonical)


def register_author_alias(
    library_root: Path, alias: str, canonical_name: str, decided_by: str, reason: str, orcid: str = "",
) -> dict[str, str]:
    alias, canonical_name, decided_by, reason = (str(value).strip() for value in (alias, canonical_name, decided_by, reason))
    if not all((alias, canonical_name, decided_by, reason)):
        raise ValueError("alias, canonical name, decider and reason are required")
    root = library_root.resolve()
    with connect_library(root) as connection:
        for shown in (alias, canonical_name):
            connection.execute(
                """INSERT INTO author_aliases(alias_normalized,alias,canonical_name,orcid,decided_by,decision_reason,updated_at)
                   VALUES (?,?,?,?,?,?,?) ON CONFLICT(alias_normalized) DO UPDATE SET alias=excluded.alias,
                   canonical_name=excluded.canonical_name,orcid=excluded.orcid,decided_by=excluded.decided_by,
                   decision_reason=excluded.decision_reason,updated_at=excluded.updated_at""",
                (_author_key(shown), shown, canonical_name, orcid, decided_by, reason, utc_now()),
            )
        affected = [row["work_id"] for row in connection.execute("SELECT work_id,author FROM works").fetchall()
                    if any(_author_key(name) in {_author_key(alias), _author_key(canonical_name)} for name in _author_names(row["author"]))]
        for work_id in affected:
            _sync_work_graph(connection, work_id)
    return {"alias": alias, "canonical_name": canonical_name, "orcid": orcid, "decided_by": decided_by}


def _sync_work_graph(connection: Any, work_id: str) -> None:
    work = connection.execute("SELECT * FROM works WHERE work_id = ?", (work_id,)).fetchone()
    if work is None:
        return
    edition = connection.execute(
        "SELECT publisher, publication_year FROM editions WHERE work_id = ? ORDER BY created_at LIMIT 1", (work_id,)
    ).fetchone()
    tags = [row[0] for row in connection.execute(
        "SELECT t.name FROM tags t JOIN work_tags wt ON wt.tag_id = t.tag_id WHERE wt.work_id = ?", (work_id,)
    )]
    source = _graph_node(connection, "work", work["canonical_title"], work_id)
    connection.execute("DELETE FROM knowledge_edges WHERE work_id = ? AND origin = 'bibliographic_metadata'", (work_id,))
    relations: list[tuple[str, str, str]] = []
    if work["author"].strip():
        for author in _author_names(work["author"]):
            canonical, normalized = _author_identity(connection, author)
            relations.append((source, "authored_by", _graph_node(connection, "person", canonical, normalized)))
    if edition and edition["publication_year"].strip():
        relations.append((source, "published_in_year", _graph_node(connection, "year", edition["publication_year"])))
    if edition and edition["publisher"].strip():
        relations.append((source, "published_by", _graph_node(connection, "organization", edition["publisher"])))
    relations.append((source, "material_type", _graph_node(connection, "material_type", work["material_type"])))
    for tag in tags:
        relation = "shelved_as" if tag.startswith("shelf:") else "tagged_as"
        relations.append((source, relation, _graph_node(connection, "tag", tag)))
    for source_id, relation, target_id in relations:
        edge_id = "KGE_" + hashlib.sha256(f"{source_id}\0{relation}\0{target_id}\0{work_id}".encode("utf-8")).hexdigest()[:24]
        connection.execute(
            """INSERT OR IGNORE INTO knowledge_edges(
                   edge_id, source_node_id, relation, target_node_id, work_id, origin, created_at
               ) VALUES (?, ?, ?, ?, ?, 'bibliographic_metadata', ?)""",
            (edge_id, source_id, relation, target_id, work_id, utc_now()),
        )


def _clean_title(value: str, fallback: str) -> str:
    title = " ".join(value.replace("\x00", " ").split()).strip(" -_")
    return title[:300] or fallback


def _filename_bibliography(path: Path) -> tuple[str, str, str]:
    """Return a conservative title/year/publisher suggestion from the file name."""
    raw = html.unescape(path.stem).replace("\u3000", " ")
    raw = re.sub(r"^\d{6,8}[\s_-]+(?=[\u4e00-\u9fff])", "", raw)
    raw = re.sub(r"^\d{8,}(?=[A-Za-z])", "", raw)
    raw = re.sub(r"^\d{1,2}[.\s_-]+(?=[\u4e00-\u9fff《])", "", raw)
    raw = re.sub(r"[_\s-]+20\d{10,}$", "", raw)
    raw = re.sub(r"[_\s-]+B0[A-Z0-9]{7,}$", "", raw, flags=re.IGNORECASE)
    raw = raw.replace("_", " ")
    raw = re.sub(r"\s*[（(](?:Z[- ]?Library|张莉(?:批注|勾画)|批注版|无水印)[^）)]*[）)]\s*", " ", raw, flags=re.IGNORECASE)
    raw = re.sub(r"(?:\s*[（(]\d+[）)])+$", "", raw)
    raw = re.sub(r"(?:\s+|[-_])*(?:[0-9]{7,})$", "", raw)
    title = _clean_title(raw, path.stem)

    year = ""
    publisher = ""
    source = re.search(
        r"[（(](?:原载于)?\s*([^（）()]{2,40}?)(1[89][0-9]{2}|20[0-9]{2})年(?:[0-9]{1,2}月|第[0-9]+期)",
        path.stem,
    )
    if source:
        publisher = _clean_title(source.group(1), "")
        publisher = re.sub(r"^(?:出版商|出版者)\s*[:：]\s*", "", publisher)
        year = source.group(2)
    if not year:
        dated_issue = re.search(r"(?<![-—–])(?<!\d)(1[89][0-9]{2}|20[0-9]{2})年\s*(?:第[0-9]+期|[0-9]{1,2}月)", path.stem)
        year = dated_issue.group(1) if dated_issue else ""
    return title, year, publisher


def _filename_is_identifier(title: str) -> bool:
    compact = re.sub(r"[^0-9A-Za-z]", "", title)
    return bool(
        re.fullmatch(r"[0-9a-fA-F]{24,}", compact)
        or re.fullmatch(r"[0-9.]+v?[0-9]*", title, re.IGNORECASE)
        or (len(compact) >= 8 and not re.search(r"[\u4e00-\u9fff]", title) and not re.search(r"[A-Za-z]{5,}", title))
    )


def _copy_suffix_rank(value: str | Path) -> tuple[bool, str]:
    path = Path(value)
    return bool(re.search(r"[（(]\d+[）)]$", path.stem)), str(path).casefold()


def _clean_author(value: str) -> str:
    author = _clean_title(str(value or ""), "")
    folded = author.casefold()
    invalid = {"cnki", "administrator", "unknown", "author", "anonymous", "n/a"}
    if folded in invalid or not re.search(r"[A-Za-zА-Яа-яЁё\u4e00-\u9fff]", author):
        return ""
    return author


def _deduplication_key(title: str) -> str:
    normalized = " ".join(str(title).split()).casefold()
    generic = {"file", "document", "untitled", "附件", "读书报告", "国际", "rp", "mmexport"}
    return "" if len(normalized) < 4 or normalized in generic else normalized


def _bibliographic_identifiers(text: str) -> set[str]:
    """Return stable work identifiers visible in title/copyright-page text."""
    values: set[str] = set()
    for match in re.findall(r"(?i)\b10\.\d{4,9}/[-._;()/:a-z0-9]+", text or ""):
        doi = match.rstrip(".,;:)]}>").casefold()
        if len(doi) >= 8:
            values.add(f"doi:{doi}")
    for match in re.finditer(r"(?i)(?:e-?isbn|isbn(?:\s*电子|\s*electronic)?)\s*[:：]?\s*([0-9xX][0-9xX\s-]{8,24})", text or ""):
        isbn = re.sub(r"[^0-9X]", "", match.group(1).upper())
        if len(isbn) in {10, 13}:
            values.add(f"isbn:{isbn}")
    return {value for value in values if not (value.startswith("doi:") and any(other != value and other.startswith(value) for other in values))}


def _year(text: str) -> str:
    match = re.search(r"(?<!\d)(1[0-9]{3}|20[0-9]{2})(?!\d)", text)
    return match.group(1) if match else ""


def _language(text: str) -> str:
    if re.search(r"[\u4e00-\u9fff]", text):
        return "zh"
    return "und"


def _material_type(title: str, sample: str, path: str = "") -> str:
    combined = f"{path}\n{title}\n{sample[:5000]}".lower()
    if any(term in combined for term in ("读书报告", "读书笔记", "读书报", "阅读札记")) or ("课程" in combined and "作业" in combined):
        return "reading_note"
    if any(term in combined for term in (
        "个人论文与稿件", "我的文章", "我的论文", "手稿", "草稿", "返修稿", "投稿稿", "未刊稿",
    )):
        return "personal_manuscript"
    if any(term in combined for term in ("博士学位", "硕士学位", "学位论文", "dissertation", "thesis")):
        return "thesis"
    if any(term in combined for term in (
        "journal", "doi", "期刊", "学报", "研究论文", "学术论文", "文章编号", "收稿日期",
    )) and not any(term in combined for term in ("books.", "isbn", "出版商", "année d'édition", "nombre de pages")):
        return "article"
    if any(term in combined for term in (
        "isbn", "cip", "出版社", "出版商", "éditeur", "année d'édition", "nombre de pages", "学术专著", "monograph",
    )):
        return "monograph"
    if any(term in combined for term in (
        "档案", "奏折", "公文", "日记", "书信", "地方志", "史料汇编", "archive", "diary",
    )):
        return "archival_source"
    if any(term in combined for term in (
        "工具书", "目录", "索引", "年鉴", "辞典", "百科", "手册", "bibliography", "catalog", "encyclopedia", "handbook",
    )):
        return "reference_work"
    return "book_or_document"


def _pdf_bibliography(
    sample: str,
    fallback_title: str,
    author: str,
    publisher: str,
    year: str,
    path: Path | None = None,
    material_type: str = "",
) -> tuple[str, str, str, str]:
    """Read bibliographic facts from title, first, and copyright pages.

    File-name and tag hints are fallbacks.  Page text has priority so journal
    articles do not remain anonymous merely because their author line lacks
    labels such as ``著`` or ``编``.
    """
    title = fallback_title
    lines = [" ".join(line.replace("\u3000", " ").split()) for line in sample[:50000].splitlines()]
    lines = [line for line in lines if line]
    compact_sample = re.sub(r"\s+", "", sample[:12000])

    def clean_people(value: str) -> str:
        value = re.sub(r"(?:作者简介|第一作者简介|通讯作者)\s*[:：]?", "", value)
        value = re.sub(r"[0-9*＊#]+", "", value)
        value = re.sub(r"\s*[（(](?:dir\.?|director|directeur|ed\.?|editor|导演|主编)[^）)]*[）)]", "", value, flags=re.IGNORECASE)
        value = re.split(r"(?:E-?mail|邮箱|收稿日期|摘要|Abstract)", value, maxsplit=1, flags=re.IGNORECASE)[0]
        parts = [part.strip(" ,，、;；·") for part in re.split(r"[,，、;；]+|\s{2,}|(?<=[A-Za-z])\s*(?:and|et|und|和)\s*(?=[A-Z])", value, flags=re.IGNORECASE)]
        people: list[str] = []
        for part in parts:
            if any(term in part for term in (
                "会议", "手册", "报告", "论文", "研究", "大学", "学院", "期刊", "学报",
                "方向", "好处", "坏处", "内容", "问题", "数据", "方案", "方法", "因此", "并且", "目前", "可以", "这个", "一种", "主要", "如果", "以及",
            )):
                continue
            if re.fullmatch(r"[\u4e00-\u9fff]{2,6}", part):
                people.append(part)
            elif re.fullmatch(r"[A-Z][A-Za-z'’-]+(?:\s+(?:[A-Z]\.|[A-Z][A-Za-z'’-]+|[A-Z]{2,})){1,5}", part):
                people.append(part)
            elif re.fullmatch(r"[А-ЯЁ](?:\.|[А-Яа-яЁё'’-]+)(?:\s+(?:[А-ЯЁ]\.|[А-ЯЁ][А-Яа-яЁё'’-]+)){1,5}", part):
                people.append(part.rstrip("."))
        return "；".join(dict.fromkeys(people))

    citation_author = citation_journal = citation_year = ""
    normalized = " ".join(lines[:120])
    chinese_citation = re.search(
        r"中文引用格式\s*[:：]\s*(?P<author>.{2,120}?)[。.]\s*"
        r"(?P<year>(?:19|20)\d{2})[。.]\s*(?P<title>.{2,300}?)[。.]\s*"
        r"(?P<journal>[^,，。]{2,80})[,，]",
        normalized,
    )
    if chinese_citation:
        citation_author = clean_people(chinese_citation.group("author"))
        citation_journal = _clean_title(chinese_citation.group("journal"), "")
        citation_year = chinese_citation.group("year")
    english_citation = re.search(
        r"(?P<author>[A-Z][A-Za-z'’. -]{2,80}),\s*[“\"](.{2,260}?)[”\"],\s*"
        r"(?P<journal>[A-Z][A-Za-z &:'’.-]{2,80})\s+\d+\s*\((?P<year>(?:19|20)\d{2})\)",
        normalized,
    )
    if english_citation and not citation_author:
        citation_author = clean_people(english_citation.group("author"))
        citation_journal = _clean_title(english_citation.group("journal"), "")
        citation_year = english_citation.group("year")
    if not citation_journal:
        simple_english_issue = re.search(
            r"\b(Environmental History|Science China Earth Sciences)\s+\d+\s*\(((?:19|20)\d{2})\)",
            normalized,
            re.IGNORECASE,
        )
        if simple_english_issue:
            citation_journal = _clean_title(simple_english_issue.group(1), "")
            citation_year = simple_english_issue.group(2)

    confirmed_filename_author = ""
    if path is not None:
        raw_stem = re.sub(r"(?:\s*[（(]\d+[）)])+$", "", html.unescape(path.stem)).strip()
        candidates: list[str] = []
        if "_" in raw_stem:
            candidates.append(raw_stem.rsplit("_", 1)[-1])
        suffix = re.search(r"(?:\s|_)([\u4e00-\u9fff]{2,6})$", raw_stem)
        if suffix:
            candidates.append(suffix.group(1))
        for candidate in candidates:
            candidate = _clean_title(candidate, "")
            if candidate and re.sub(r"\s+", "", candidate) in compact_sample:
                confirmed_filename_author = clean_people(candidate)
                if confirmed_filename_author:
                    break

    named = re.search(r"书\s*名\s*[:：]?\s*([^\n]{2,60})", sample[:8000])
    if named:
        title = named.group(1).strip()
    role_line = next((index for index, line in enumerate(lines[:12]) if re.search(r"[（(](?:dir\.?|director|directeur|ed\.?|editor|导演|主编)[^）)]*[）)]", line, re.IGNORECASE)), -1)
    if role_line > 0 and any(token in sample[:3000] for token in ("DOI", "ISBN", "Éditeur", "出版商", "出版社")):
        page_title = _clean_title("：".join(lines[:role_line]), "")
        if 5 <= len(page_title) <= 300 and (fallback_title.startswith("(NEW)") or _filename_is_identifier(fallback_title)):
            title = page_title
    page_author = citation_author or confirmed_filename_author
    if not page_author:
        def biblio_key(value: str) -> str:
            return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", value.casefold())

        compact_title = biblio_key(fallback_title)
        title_index = -1
        title_span_end = -1
        title_probe = compact_title[: min(80, len(compact_title))]
        for index in range(min(80, len(lines))):
            for width in range(1, 5):
                window = biblio_key(" ".join(lines[index:index + width]))
                if len(title_probe) >= 6 and (title_probe in window or (len(window) >= 12 and window in compact_title)):
                    title_index, title_span_end = index, index + width
                    break
            if title_index >= 0:
                break
        nearby: list[str] = []
        if title_index >= 0 and material_type == "article":
            nearby.extend(lines[max(0, title_index - 1):title_index])
            nearby.extend(lines[title_span_end:title_span_end + 5])
        for line in nearby:
            if any(term in line for term in ("摘要", "关键词", "大学", "学院", "研究所", "学报", "期刊", "DOI", "卷", "期")):
                continue
            candidate = clean_people(line)
            if candidate:
                page_author = candidate
                break
        if not page_author and lines and material_type == "article":
            first_candidate = clean_people(lines[0])
            following = biblio_key(" ".join(lines[1:5]))
            if first_candidate and title_probe and (title_probe in following or following in compact_title):
                page_author = first_candidate
        if not page_author and material_type == "article":
            for index, line in enumerate(lines[:6]):
                candidate = clean_people(line)
                following = " ".join(lines[index + 1:index + 5])
                if candidate and re.search(r"(?:摘要|Abstract|学报|期刊|Journal|\d{4}年第?\d+期)", following, re.IGNORECASE):
                    page_author = candidate
                    break
    if not page_author:
        authored = re.search(r"^([^\n]{2,40}(?:著|撰|编|校点))\s*$", sample[:4000], re.MULTILINE)
        page_author = authored.group(1).strip() if authored else ""
    if not page_author:
        directed = re.search(
            r"^([^\n]{3,160}?)(?:\s*[（(](?:dir\.?|director|directeur|ed\.?|editor|导演|主编)[^）)]*[）)])\s*$",
            sample[:4000], re.MULTILINE | re.IGNORECASE,
        )
        page_author = clean_people(directed.group(1)) if directed else ""
    if not page_author:
        for index, line in enumerate(lines[:20]):
            candidate = clean_people(line.strip(" ."))
            preceding = " ".join(lines[max(0, index - 2):index]).casefold()
            if candidate and re.search(r"[А-ЯЁа-яё]", candidate) and re.search(r"(?:автор|дневник|экспедици|под\s+ред)", preceding):
                page_author = candidate
                break

    page_publisher = ""
    if material_type == "article" or citation_journal:
        page_publisher = citation_journal
        if not page_publisher:
            for line in lines[:50]:
                if len(line) > 90 or any(mark in line for mark in ("《", "》", "；", ";")):
                    continue
                match = re.search(
                    r"(?:[\u4e00-\u9fffA-Za-z ]{0,45}(?:学报(?:\s*[（(][^）)]{1,30}[）)])?|论丛|馆刊|农业考古)"
                    r"|Journal of [A-Za-z &:'’.-]{3,80}|Environmental History|Science China Earth Sciences)",
                    line,
                    re.IGNORECASE,
                )
                if match:
                    page_publisher = _clean_title(match.group(0), "")
                    break
            if not page_publisher and "中国国家博物馆馆刊" in re.sub(r"\s+", "", sample[:5000]):
                page_publisher = "中国国家博物馆馆刊"
    if not page_publisher and material_type != "article":
        published = re.search(r"^(?:出版商|出版者|出版社|Éditeur)\s*[:：]\s*([^\n]{2,100})$", sample[:12000], re.MULTILINE | re.IGNORECASE)
        if not published:
            published = re.search(r"^([^\n]{2,70}(?:出版社|University Press|Publishing House|Press))\s*$", sample[:12000], re.MULTILINE | re.IGNORECASE)
        page_publisher = published.group(1).strip() if published else ""
    page_publisher = re.sub(r"^(?:出版商|出版者)\s*[:：]\s*", "", page_publisher).strip(" ,，、;；:：")
    if page_publisher.casefold() == "environmental history":
        page_publisher = "Environmental History"

    page_year = citation_year
    if not page_year and material_type == "article" and page_publisher:
        issue_date = re.search(
            r"(?<!\d)((?:19|20)\d{2})\s*年\s*(?:第\s*\d+\s*期|\d+\s*月)",
            "\n".join(lines[:25]),
        )
        page_year = issue_date.group(1) if issue_date else ""
        if not page_year:
            journal_date = re.search(
                rf"{re.escape(page_publisher)}[^\n]{{0,30}}(?<!\d)((?:19|20)\d{{2}})(?!\d)",
                "\n".join(lines[:30]),
                re.IGNORECASE,
            )
            page_year = journal_date.group(1) if journal_date else ""
        if not page_year and page_publisher == "Environmental History":
            environmental_history_date = re.search(r"Environmental History\s+\d+\s*\(((?:19|20)\d{2})\)", normalized, re.IGNORECASE)
            page_year = environmental_history_date.group(1) if environmental_history_date else ""
    if not page_year:
        dated = re.search(r"(?:出版日期|出版年份|Année d['’]édition|版次|CIP)[^\n]{0,50}(1[0-9]{3}|20[0-9]{2})", sample[:3000], re.IGNORECASE)
        page_year = dated.group(1) if dated else ""

    return title, page_author or author, page_publisher or publisher, page_year or year


def _triage_state(sample: str, supported: bool, text_layer: str, page_count: int | None) -> tuple[str, str]:
    if not supported:
        return "unsupported", "当前版本暂不解析此格式；文件仍保留在盘点清单中。CAJ 应保留原件并另生成带转换回执的 PDF/逐页 Markdown。"
    if page_count and text_layer == "absent":
        return "needs_visual_triage", "前十页没有可用文本层，需要原页或视觉模型进一步判断。"
    normalized = sample.lower()
    score = sum(1 for term in HISTORY_TERMS if term in normalized)
    if score >= 2:
        return "likely_historical", f"前段文本命中 {score} 个历史材料线索，仅作为整理建议。"
    if len(sample.strip()) < 100:
        return "uncertain", "可读取文本过少，不能可靠判断材料类型。"
    if score == 0 and len(sample) >= 500:
        return "not_obviously_historical", "前段文本未出现明显历史材料线索，不代表不可用于研究。"
    return "uncertain", "有少量线索但不足以稳定判断，建议人工确认。"


def _word_intake_decision(path: Path, item: dict[str, Any]) -> tuple[str, str]:
    if path.suffix.lower() != ".docx":
        return "admit", ""
    stem = re.sub(r"(?:\s*[（(]\d+[）)])+$", "", path.stem).strip()
    if stem.startswith("~$"):
        return "ignore", "Office 临时锁文件"
    negative = (
        "申请表", "审批表", "回执", "合同", "简历", "模板", "住宿晚归", "学籍在线验证",
        "操作手册", "指导手册", "练习指导", "实习教程", "课程实习", "剧本", "小品", "mod介绍",
        "报告表", "会议邀请", "选择题", "讲稿",
    )
    if any(term.casefold() in stem.casefold() for term in negative):
        return "ignore", "行政、教学或非研究 Word"
    if re.fullmatch(r"(?:摘要|引言|开头|小结论|编者前言|框架\d*(?:（[^）]*）)?)(?:\s*[-—_].*)?", stem, re.IGNORECASE):
        return "ignore", "孤立的写作碎片"
    sample = str(item.get("sample_text", ""))
    if "示例稿" in sample[:1000] and "用于演示" in sample[:3000]:
        return "ignore", "演示文件"
    material_type = str(item.get("suggested_material_type", ""))
    admit_terms = (
        "翻译", "译文", "译稿", "原作", "史料", "档案", "日记", "游记", "考察", "材料", "资料整理",
        "地方志",
    )
    if material_type == "archival_source" or any(term.casefold() in stem.casefold() for term in admit_terms):
        return "admit", ""
    if re.search(r"(?:дневник|путешествие|экспедици)", sample[:3000], re.IGNORECASE):
        return "admit", ""
    review_types = {"article", "monograph", "thesis", "personal_manuscript", "reading_note"}
    review_terms = ("论文", "研究", "专著", "读书报告", "读书笔记", "环境史", "行政区划", "气候", "地理")
    if material_type in review_types or any(term.casefold() in stem.casefold() for term in review_terms):
        return "review", "需先比较正文完整度、重复内容和版本关系。"
    if len(sample) >= 5000 and int(item.get("byte_count", 0) or 0) >= 20000:
        return "review", "长篇 Word 需先核对是否为完整研究稿及当前版本。"
    return "ignore", "未见完整研究材料特征"


def _inspect_file(path: Path) -> dict[str, Any]:
    stat = path.stat()
    suffix = path.suffix.lower()
    supported = suffix in SUPPORTED_SUFFIXES
    title, filename_year, filename_publisher = _filename_bibliography(path)
    author = ""
    publisher = ""
    year = ""
    page_count: int | None = None
    inspected_pages = 0
    text_layer = "not_applicable"
    sample = ""

    if suffix == ".pdf":
        with fitz.open(path) as document:
            page_count = document.page_count
            metadata = document.metadata or {}
            embedded_title = _clean_title(str(metadata.get("title") or ""), "")
            if embedded_title and _filename_is_identifier(title):
                title = embedded_title
            embedded_author = metadata.get("author") or ""
            chunks: list[str] = []
            for index in range(min(10, page_count)):
                chunks.append(document.load_page(index).get_text("text"))
                inspected_pages += 1
            sample = "\n".join(chunks).strip()[:50000]
            text_layer = "present" if sample else "absent"
            material_type = _material_type(title, sample, str(path))
            page_title, page_author, page_publisher, page_year = _pdf_bibliography(
                sample, title, "", "", "", path=path, material_type=material_type,
            )
            title = page_title or title
            author = page_author or embedded_author
            publisher = page_publisher or filename_publisher
            year = page_year or filename_year
    elif suffix in {".md", ".txt", ".csv", ".tsv"}:
        sample = path.read_text(encoding="utf-8", errors="replace")[:50000]
        inspected_pages = 1
        text_layer = "present" if sample.strip() else "absent"
    elif suffix == ".docx":
        from docx import Document

        document = Document(path)
        chunks = [" ".join(paragraph.text.split()) for paragraph in document.paragraphs if paragraph.text.strip()]
        sample = "\n".join(chunks)[:50000]
        inspected_pages = 1
        text_layer = "present" if sample.strip() else "absent"
    elif suffix in {".xlsx", ".xlsm"}:
        from .attachments import _xlsx_preview

        previews = _xlsx_preview(path)
        sample = "\n".join(
            "\t".join(str(value) for value in row)
            for sheet in previews for row in sheet.get("rows", [])[:20]
        )[:50000]
        if previews and previews[0].get("sheet"):
            title = _clean_title(path.stem, previews[0]["sheet"])
        inspected_pages = len(previews) or 1
        text_layer = "present" if sample.strip() else "absent"
    elif suffix in {".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff"}:
        inspected_pages = 1
        text_layer = "absent"

    title = _clean_title(title, path.stem)
    material_type = _material_type(title, sample, str(path))
    if sample and suffix in {".md", ".txt", ".docx"}:
        _, page_author, page_publisher, page_year = _pdf_bibliography(
            sample, title, "", "", "", path=path, material_type=material_type,
        )
        author = page_author or author
        publisher = page_publisher or filename_publisher
        year = page_year or filename_year
    first_author = re.split(r"[；;,，、]", _clean_author(author))[0].strip()
    if first_author and title.endswith(f" {first_author}") and re.sub(r"\s+", "", first_author) in re.sub(r"\s+", "", sample[:12000]):
        title = title[: -(len(first_author) + 1)].rstrip()
    state, reason = _triage_state(sample, supported, text_layer, page_count)
    return {
        "path": str(path.resolve()),
        "format": suffix.lstrip(".") or "unknown",
        "byte_count": stat.st_size,
        "modified_ns": stat.st_mtime_ns,
        "sha256": _file_hash(path),
        "suggested_title": title,
        "suggested_author": _clean_author(author),
        "suggested_year": year,
        "suggested_publisher": _clean_title(publisher, "") if publisher else "",
        "suggested_language": _language(f"{title}\n{sample[:4000]}"),
        "suggested_material_type": material_type,
        "page_count": page_count,
        "text_layer": text_layer,
        "triage_state": state,
        "triage_reason": reason,
        "inspected_pages": inspected_pages,
        "sample_text": sample,
    }


def library_root_for(project_root: Path, library_root: Path | None = None) -> Path:
    root = resolve_library_root(project_root, library_root)
    initialize_library(root)
    return root


def library_status(project_root: Path, library_root: Path | None = None) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        counts = {
            table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in ("works", "editions", "library_files", "file_versions", "scan_sessions")
        }
    return {"library_root": str(root), "counts": counts, "skills": discover_skills()}


def _candidate_action(connection: Any, item: dict[str, Any]) -> dict[str, Any]:
    existing_file = connection.execute(
        "SELECT file_id, work_id, edition_id FROM library_files WHERE path = ?",
        (item["path"],),
    ).fetchone()
    if existing_file is not None:
        current = connection.execute(
            "SELECT sha256 FROM file_versions WHERE file_id = ? AND is_current = 1",
            (existing_file["file_id"],),
        ).fetchone()
        action = "unchanged" if current and current["sha256"] == item["sha256"] else "new_version"
        return {"proposed_action": action, **dict(existing_file)}
    duplicate = connection.execute(
        """SELECT f.file_id, f.work_id, f.edition_id
           FROM file_versions v JOIN library_files f ON f.file_id = v.file_id
           WHERE v.sha256 = ? LIMIT 1""",
        (item["sha256"],),
    ).fetchone()
    if duplicate is not None:
        return {"proposed_action": "exact_duplicate", **dict(duplicate)}
    same_work = _existing_work_for_candidate(connection, item)
    if same_work is not None:
        return {"proposed_action": "same_work", **dict(same_work), "file_id": None}
    return {
        "proposed_action": "register_new",
        "file_id": None,
        "work_id": None,
        "edition_id": None,
    }


def _existing_work_for_title(connection: Any, title: str) -> Any:
    key = _deduplication_key(title)
    if not key:
        return None
    return connection.execute(
        """SELECT w.work_id, e.edition_id
           FROM works w JOIN editions e ON e.work_id = w.work_id
           WHERE w.canonical_title = ? COLLATE NOCASE
           ORDER BY w.updated_at DESC, e.created_at LIMIT 1""",
        (" ".join(str(title).split()),),
    ).fetchone()


def _existing_work_for_candidate(connection: Any, candidate: Any) -> Any:
    for identifier in sorted(_bibliographic_identifiers(str(candidate["sample_text"] or ""))):
        kind, value = identifier.split(":", 1)
        if kind == "doi":
            row = connection.execute(
                """SELECT f.work_id, f.edition_id FROM file_versions v
                   JOIN library_files f ON f.file_id=v.file_id
                   WHERE v.is_current=1 AND instr(lower(v.sample_text), ?) > 0
                   ORDER BY v.discovered_at DESC LIMIT 1""",
                (value,),
            ).fetchone()
        else:
            row = connection.execute(
                """SELECT f.work_id, f.edition_id FROM file_versions v
                   JOIN library_files f ON f.file_id=v.file_id
                   WHERE v.is_current=1 AND instr(
                       replace(replace(replace(replace(upper(v.sample_text), '-', ''), ' ', ''), char(10), ''), char(13), ''), ?
                   ) > 0 ORDER BY v.discovered_at DESC LIMIT 1""",
                (value,),
            ).fetchone()
        if row is not None:
            return row
    return _existing_work_for_title(connection, candidate["suggested_title"])


def create_scan_session(
    project_root: Path,
    source_root: Path,
    library_root: Path | None = None,
    skill_name: str = "historical-material-intake",
) -> dict[str, Any]:
    source_root = source_root.expanduser().resolve()
    if not source_root.is_dir():
        raise FileNotFoundError(f"scan directory does not exist: {source_root}")
    skill = get_skill(skill_name)
    if "library_intake" not in skill["compatible_actions"]:
        raise ValueError(f"skill is not compatible with library intake: {skill_name}")
    root = library_root_for(project_root, library_root)
    session_id = f"SCN_{uuid.uuid4().hex}"
    now = utc_now()
    with connect_library(root) as connection:
        connection.execute(
            """INSERT INTO scan_sessions(
                   session_id, root_path, skill_name, skill_sha256, status, created_at
               ) VALUES (?, ?, ?, ?, 'scanning', ?)""",
            (session_id, str(source_root), skill["name"], skill["sha256"], now),
        )
    return scan_session(project_root, session_id, root)


def run_scan_session(
    project_root: Path,
    session_id: str,
    library_root: Path | None = None,
) -> None:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        session = connection.execute(
            "SELECT * FROM scan_sessions WHERE session_id = ?", (session_id,)
        ).fetchone()
        if session is None:
            raise KeyError(f"unknown scan session: {session_id}")
        source_root = Path(session["root_path"])
    processed = 0
    pending: list[tuple[str, dict[str, Any], str, str, dict[str, Any]]] = []
    seen_identifiers: set[str] = set()

    def flush_pending() -> None:
        nonlocal processed
        if not pending:
            return
        with connect_library(root) as connection:
            for candidate_id, item, status, error, fallback_identity in pending:
                identity = (
                    _candidate_action(connection, item) if status not in {"error", "ignored"} else fallback_identity
                )
                identifiers = _bibliographic_identifiers(str(item.get("sample_text", "")))
                if identity["proposed_action"] == "register_new" and identifiers & seen_identifiers:
                    identity = {"proposed_action": "same_scan_work", "work_id": None, "edition_id": None, "file_id": None}
                seen_identifiers.update(identifiers)
                connection.execute(
                    """INSERT INTO scan_candidates(
                           candidate_id, session_id, path, format, byte_count, modified_ns, sha256,
                           suggested_title, suggested_author, suggested_year, suggested_publisher,
                           suggested_language, suggested_material_type, page_count, text_layer,
                           triage_state, triage_reason, inspected_pages, sample_text, proposed_action,
                           existing_work_id, existing_edition_id, existing_file_id, status, error
                       ) VALUES (
                           ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                       )""",
                    (
                        candidate_id, session_id, item["path"], item["format"], item["byte_count"],
                        item["modified_ns"], item["sha256"], item["suggested_title"],
                        item["suggested_author"], item["suggested_year"], item["suggested_publisher"],
                        item["suggested_language"], item["suggested_material_type"], item["page_count"],
                        item["text_layer"], item["triage_state"], item["triage_reason"],
                        item["inspected_pages"], item["sample_text"], identity["proposed_action"],
                        identity.get("work_id"), identity.get("edition_id"), identity.get("file_id"),
                        status, error,
                    ),
                )
            processed += len(pending)
            connection.execute(
                "UPDATE scan_sessions SET processed_count = ? WHERE session_id = ?",
                (processed, session_id),
            )
        pending.clear()

    try:
        for directory, _, filenames in os.walk(source_root, followlinks=False):
            for filename in sorted(filenames, key=_copy_suffix_rank):
                path = Path(directory) / filename
                if path.suffix.lower() not in CANDIDATE_SUFFIXES:
                    continue
                if _automatic_scan_exclusion(path, source_root):
                    continue
                candidate_id = f"CND_{uuid.uuid4().hex}"
                if path.suffix.lower() == ".docx" and path.name.startswith("~$"):
                    stat = path.stat()
                    item = {
                        "path": str(path.resolve()), "format": "docx", "byte_count": stat.st_size,
                        "modified_ns": stat.st_mtime_ns, "sha256": "", "suggested_title": path.stem,
                        "suggested_author": "", "suggested_year": "", "suggested_publisher": "",
                        "suggested_language": "und", "suggested_material_type": "unknown",
                        "page_count": None, "text_layer": "unknown", "triage_state": "ignored_word",
                        "triage_reason": "Office 临时锁文件", "inspected_pages": 0, "sample_text": "",
                    }
                    pending.append((candidate_id, item, "ignored", "", {"proposed_action": "ignored", "work_id": None, "edition_id": None, "file_id": None}))
                    if len(pending) >= SCAN_WRITE_BATCH_SIZE:
                        flush_pending()
                    continue
                try:
                    item = _inspect_file(path)
                    identity = {"proposed_action": "", "work_id": None, "edition_id": None, "file_id": None}
                    route, route_reason = _word_intake_decision(path, item)
                    status = "ignored" if route == "ignore" else "preview"
                    if route == "ignore":
                        item["triage_state"] = "ignored_word"
                        item["triage_reason"] = route_reason
                        identity["proposed_action"] = "ignored"
                    elif route == "review":
                        item["triage_state"] = "word_review"
                        item["triage_reason"] = route_reason
                    error = ""
                except Exception as exc:
                    stat = path.stat()
                    item = {
                        "path": str(path.resolve()), "format": path.suffix.lower().lstrip("."),
                        "byte_count": stat.st_size, "modified_ns": stat.st_mtime_ns, "sha256": "",
                        "suggested_title": path.stem, "suggested_author": "", "suggested_year": "",
                        "suggested_publisher": "", "suggested_language": "und",
                        "suggested_material_type": "unknown", "page_count": None,
                        "text_layer": "unknown", "triage_state": "error",
                        "triage_reason": "文件无法完成只读盘点。", "inspected_pages": 0, "sample_text": "",
                    }
                    identity = {"proposed_action": "error", "work_id": None, "edition_id": None, "file_id": None}
                    status = "error"
                    error = str(exc)
                pending.append((candidate_id, item, status, error, identity))
                if len(pending) >= SCAN_WRITE_BATCH_SIZE:
                    flush_pending()
        flush_pending()
        with connect_library(root) as connection:
            connection.execute(
                """UPDATE scan_sessions
                   SET status = 'preview_ready', processed_count = ?, completed_at = ?
                   WHERE session_id = ?""",
                (processed, utc_now(), session_id),
            )
    except Exception as exc:
        pending.clear()
        with connect_library(root) as connection:
            connection.execute(
                """UPDATE scan_sessions
                   SET status = 'failed', error = ?, processed_count = ?, completed_at = ?
                   WHERE session_id = ?""",
                (str(exc), processed, utc_now(), session_id),
            )


def start_scan_session(
    project_root: Path,
    source_root: Path,
    library_root: Path | None = None,
    skill_name: str = "historical-material-intake",
    compatibility_wait: float = 0.0,
) -> dict[str, Any]:
    session = create_scan_session(project_root, source_root, library_root, skill_name)
    root = library_root_for(project_root, library_root)
    worker = threading.Thread(
        target=run_scan_session,
        args=(project_root, session["session_id"], root),
        name=f"library-scan-{session['session_id']}",
        daemon=True,
    )
    worker.start()
    # Optional only for callers that explicitly want a brief compatibility wait; the
    # HTTP endpoint uses the non-blocking default.
    if compatibility_wait > 0:
        worker.join(timeout=compatibility_wait)
        return scan_session(project_root, session["session_id"], root)
    # The session row was durably committed before the worker started. Returning
    # that receipt avoids contending with the worker's first SQLite transaction
    # and keeps the HTTP start endpoint genuinely non-blocking.
    return session


def scan_directory(
    project_root: Path,
    source_root: Path,
    library_root: Path | None = None,
    skill_name: str = "historical-material-intake",
) -> dict[str, Any]:
    """Synchronous compatibility helper used by the CLI and existing small-directory callers."""
    session = create_scan_session(project_root, source_root, library_root, skill_name)
    run_scan_session(project_root, session["session_id"], library_root)
    return scan_session(project_root, session["session_id"], library_root)


def archive_uploaded_file(
    project_root: Path,
    source_path: Path,
    library_root: Path | None = None,
    display_name: str = "",
) -> dict[str, Any]:
    """Register one chat-uploaded file immediately, reusing the library hash/version rules."""
    source_path = source_path.expanduser().resolve()
    session = scan_directory(project_root, source_path.parent, library_root)
    candidate = next(
        (item for item in session["candidates"] if Path(item["path"]).resolve() == source_path),
        None,
    )
    if candidate is None:
        root = library_root_for(project_root, library_root)
        with connect_library(root) as connection:
            row = connection.execute(
                "SELECT * FROM scan_candidates WHERE session_id=? AND path=?",
                (session["session_id"], str(source_path)),
            ).fetchone()
            if row is not None and row["status"] == "ignored":
                connection.execute(
                    "UPDATE scan_candidates SET status='preview',triage_state='uncertain',triage_reason=? WHERE candidate_id=?",
                    ("用户明确上传，覆盖批量 Word 准入规则。", row["candidate_id"]),
                )
                candidate = dict(row)
                candidate.update(status="preview", triage_state="uncertain")
    if candidate is None:
        raise RuntimeError("uploaded file was not found in its library intake session")
    result = approve_candidates(project_root, session["session_id"], [candidate["candidate_id"]], library_root)
    approved = result["approved"][0] if result["approved"] else {
        "work_id": candidate.get("existing_work_id", ""), "version_id": "",
    }
    work_id = str(approved.get("work_id", ""))
    title = Path(display_name).stem.strip()
    if work_id and title:
        root = library_root_for(project_root, library_root)
        with connect_library(root) as connection:
            current = connection.execute(
                "SELECT canonical_title FROM works WHERE work_id=?", (work_id,),
            ).fetchone()
            if current and (
                candidate["proposed_action"] != "unchanged"
                or
                str(current["canonical_title"]) == source_path.stem
                or re.fullmatch(r"[0-9a-f]{64}", str(current["canonical_title"]), re.IGNORECASE)
            ):
                connection.execute(
                    "UPDATE works SET canonical_title=?,updated_at=? WHERE work_id=?",
                    (title, utc_now(), work_id),
                )
                _refresh_search(connection, work_id)
                _sync_work_graph(connection, work_id)
    return {
        "session_id": session["session_id"], "candidate_id": candidate["candidate_id"],
        "action": candidate["proposed_action"], **approved,
    }


def scan_session(
    project_root: Path,
    session_id: str = "",
    library_root: Path | None = None,
    page: int = 1,
    page_size: int = SCAN_PAGE_SIZE,
) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    page = max(1, int(page))
    page_size = max(1, min(int(page_size), SCAN_MAX_PAGE_SIZE))
    with connect_library(root) as connection:
        if session_id:
            session = connection.execute(
                "SELECT * FROM scan_sessions WHERE session_id = ?", (session_id,)
            ).fetchone()
        else:
            session = connection.execute(
                "SELECT * FROM scan_sessions ORDER BY created_at DESC LIMIT 1"
            ).fetchone()
        if session is None:
            raise KeyError(f"unknown scan session: {session_id}")
        total_count = connection.execute(
            "SELECT COUNT(*) FROM scan_candidates WHERE session_id = ? AND status <> 'ignored'", (session["session_id"],)
        ).fetchone()[0]
        ignored_word_count = connection.execute(
            "SELECT COUNT(*) FROM scan_candidates WHERE session_id = ? AND status = 'ignored'", (session["session_id"],)
        ).fetchone()[0]
        word_review_count = connection.execute(
            "SELECT COUNT(*) FROM scan_candidates WHERE session_id = ? AND status='preview' AND triage_state='word_review'",
            (session["session_id"],),
        ).fetchone()[0]
        eligible_remaining_count = connection.execute(
            """SELECT COUNT(*) FROM scan_candidates
               WHERE session_id = ? AND status = 'preview'
                 AND triage_state NOT IN ('unsupported', 'error')
                 AND triage_state <> 'word_review'
                 AND proposed_action != 'unchanged'""",
            (session["session_id"],),
        ).fetchone()[0]
        offset = (page - 1) * page_size
        candidates = connection.execute(
            """SELECT c.candidate_id, c.session_id, c.path, c.format, c.byte_count, c.modified_ns, c.sha256,
                      c.suggested_title, c.suggested_author, c.suggested_year, c.suggested_publisher,
                      c.suggested_language, c.suggested_material_type, c.page_count, c.text_layer,
                      c.triage_state, c.triage_reason, c.inspected_pages, c.proposed_action,
                      c.existing_work_id, c.existing_edition_id, c.existing_file_id, c.status, c.error,
                      f.work_id AS resolved_work_id, w.canonical_title AS resolved_work_title,
                      w.author AS resolved_work_author, w.material_type AS resolved_material_type,
                      e.publisher AS resolved_publisher,
                      e.publication_year AS resolved_year,
                      (SELECT COUNT(*) FROM library_files peers WHERE peers.work_id=f.work_id) AS resolved_file_count
               FROM scan_candidates c
               LEFT JOIN library_files f ON f.path=c.path
               LEFT JOIN works w ON w.work_id=f.work_id
               LEFT JOIN editions e ON e.edition_id=f.edition_id
               WHERE c.session_id = ? AND c.status <> 'ignored' ORDER BY c.path LIMIT ? OFFSET ?""",
            (session["session_id"], page_size, offset),
        ).fetchall()
    candidate_items = []
    for row in candidates:
        item = dict(row)
        item["suggested_shelf"] = SUGGESTED_SHELVES.get(item["suggested_material_type"], "unclassified")
        item["suggested_shelf_label"] = LIBRARY_SHELVES[item["suggested_shelf"]]
        if item.get("resolved_material_type"):
            item["resolved_shelf"] = SUGGESTED_SHELVES.get(item["resolved_material_type"], "unclassified")
            item["resolved_shelf_label"] = LIBRARY_SHELVES[item["resolved_shelf"]]
        candidate_items.append(item)
    return {
        **dict(session),
        "total_count": total_count,
        "ignored_word_count": ignored_word_count,
        "word_review_count": word_review_count,
        "eligible_remaining_count": eligible_remaining_count,
        "page": page,
        "page_size": page_size,
        "page_count": (total_count + page_size - 1) // page_size,
        "has_more": offset + len(candidates) < total_count,
        "candidates": candidate_items,
    }


def _add_tag(connection: Any, work_id: str, name: str, origin: str) -> None:
    name = name.strip()
    if not name:
        return
    connection.execute("INSERT OR IGNORE INTO tags(name) VALUES (?)", (name,))
    tag_id = connection.execute("SELECT tag_id FROM tags WHERE name = ?", (name,)).fetchone()[0]
    connection.execute(
        "INSERT OR IGNORE INTO work_tags(work_id, tag_id, origin) VALUES (?, ?, ?)",
        (work_id, tag_id, origin),
    )


def _refresh_search(connection: Any, work_id: str) -> None:
    work = connection.execute("SELECT * FROM works WHERE work_id = ?", (work_id,)).fetchone()
    publishers = connection.execute(
        "SELECT GROUP_CONCAT(publisher, ' ') FROM editions WHERE work_id = ?", (work_id,)
    ).fetchone()[0] or ""
    tags = connection.execute(
        """SELECT GROUP_CONCAT(t.name, ' ') FROM tags t
           JOIN work_tags wt ON wt.tag_id = t.tag_id WHERE wt.work_id = ?""", (work_id,)
    ).fetchone()[0] or ""
    samples = connection.execute(
        """SELECT GROUP_CONCAT(substr(v.sample_text, 1, 5000), ' ') FROM file_versions v
           JOIN library_files f ON f.file_id = v.file_id WHERE f.work_id = ? AND v.is_current = 1""",
        (work_id,),
    ).fetchone()[0] or ""
    connection.execute("DELETE FROM work_search WHERE work_id = ?", (work_id,))
    connection.execute(
        "INSERT INTO work_search(work_id, title, author, publisher, tags, sample_text) VALUES (?, ?, ?, ?, ?, ?)",
        (work_id, work["canonical_title"], work["author"], publishers, tags, samples),
    )


def _insert_version(connection: Any, candidate: Any, file_id: str, skill: Any) -> str:
    version_id = f"FV_{uuid.uuid4().hex}"
    connection.execute("UPDATE file_versions SET is_current = 0 WHERE file_id = ?", (file_id,))
    connection.execute(
        """INSERT INTO file_versions(
               version_id, file_id, sha256, byte_count, modified_ns, format, page_count,
               text_layer, triage_state, triage_reason, inspected_pages, sample_text,
               qualification, skill_name, skill_sha256, discovered_at, is_current
           ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'FILE_VERIFIED', ?, ?, ?, 1)""",
        (
            version_id, file_id, candidate["sha256"], candidate["byte_count"], candidate["modified_ns"],
            candidate["format"], candidate["page_count"], candidate["text_layer"],
            candidate["triage_state"], candidate["triage_reason"], candidate["inspected_pages"],
            candidate["sample_text"], skill["name"], skill["sha256"], utc_now(),
        ),
    )
    return version_id


def approve_candidates(
    project_root: Path,
    session_id: str,
    candidate_ids: list[str] | None = None,
    library_root: Path | None = None,
) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        session = connection.execute(
            "SELECT * FROM scan_sessions WHERE session_id = ?", (session_id,)
        ).fetchone()
        if session is None:
            raise KeyError(f"unknown scan session: {session_id}")
        if session["status"] not in {"preview_ready", "partially_approved", "approved"}:
            raise ValueError("scan candidates may only be approved after the scan preview is ready")
        skill = {"name": session["skill_name"], "sha256": session["skill_sha256"]}
        parameters: list[Any] = [session_id]
        clause = ""
        if candidate_ids:
            clause = f" AND candidate_id IN ({','.join('?' for _ in candidate_ids)})"
            parameters.extend(candidate_ids)
        candidates = connection.execute(
            f"SELECT * FROM scan_candidates WHERE session_id = ?{clause} ORDER BY path", parameters
        ).fetchall()
        candidates = sorted(candidates, key=lambda candidate: _copy_suffix_rank(candidate["path"]))
        approved: list[dict[str, str]] = []
        for candidate in candidates:
            if candidate["status"] != "preview" or candidate["triage_state"] in {"error", "unsupported"}:
                continue
            if not candidate_ids and candidate["triage_state"] == "word_review":
                continue
            action = candidate["proposed_action"]
            if action in {"register_new", "same_scan_work"}:
                action = "same_work" if _existing_work_for_candidate(connection, candidate) else "register_new"
            if action == "unchanged":
                connection.execute(
                    "UPDATE library_files SET last_seen_at = ? WHERE file_id = ?",
                    (utc_now(), candidate["existing_file_id"]),
                )
                work_id = candidate["existing_work_id"]
                version_id = ""
            elif action == "new_version":
                file_id = candidate["existing_file_id"]
                work_id = candidate["existing_work_id"]
                version_id = _insert_version(connection, candidate, file_id, skill)
                connection.execute("UPDATE library_files SET last_seen_at = ? WHERE file_id = ?", (utc_now(), file_id))
            else:
                now = utc_now()
                if action in {"exact_duplicate", "same_work"}:
                    existing = (
                        _existing_work_for_candidate(connection, candidate)
                        if action == "same_work" else candidate
                    )
                    work_id = existing["work_id"] if action == "same_work" else existing["existing_work_id"]
                    edition_id = existing["edition_id"] if action == "same_work" else existing["existing_edition_id"]
                else:
                    work_id = f"WRK_{uuid.uuid4().hex}"
                    edition_id = f"ED_{uuid.uuid4().hex}"
                    connection.execute(
                        """INSERT INTO works(
                               work_id, canonical_title, author, language, material_type, created_at, updated_at
                           ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                        (
                            work_id, candidate["suggested_title"], candidate["suggested_author"],
                            candidate["suggested_language"], candidate["suggested_material_type"], now, now,
                        ),
                    )
                    connection.execute(
                        """INSERT INTO editions(
                               edition_id, work_id, edition_label, publisher, publication_year, isbn, created_at
                           ) VALUES (?, ?, '待核版本', ?, ?, '', ?)""",
                        (edition_id, work_id, candidate["suggested_publisher"], candidate["suggested_year"], now),
                    )
                file_id = f"FIL_{uuid.uuid4().hex}"
                connection.execute(
                    """INSERT INTO library_files(
                           file_id, work_id, edition_id, path, first_seen_at, last_seen_at
                       ) VALUES (?, ?, ?, ?, ?, ?)""",
                    (file_id, work_id, edition_id, candidate["path"], now, now),
                )
                version_id = _insert_version(connection, candidate, file_id, skill)
                if action == "register_new":
                    _add_tag(connection, work_id, f"material:{candidate['suggested_material_type']}", "system")
                    _add_tag(connection, work_id, f"triage:{candidate['triage_state']}", "system")
                    shelf = SUGGESTED_SHELVES.get(candidate["suggested_material_type"], "unclassified")
                    _add_tag(connection, work_id, f"shelf:{shelf}", "system")
                else:
                    work = connection.execute("SELECT author FROM works WHERE work_id=?", (work_id,)).fetchone()
                    if work and not work["author"] and candidate["suggested_author"]:
                        connection.execute(
                            "UPDATE works SET author=?,updated_at=? WHERE work_id=?",
                            (candidate["suggested_author"], utc_now(), work_id),
                        )
                    edition = connection.execute(
                        "SELECT publisher,publication_year FROM editions WHERE edition_id=?", (edition_id,)
                    ).fetchone()
                    if edition:
                        connection.execute(
                            """UPDATE editions SET publisher=?,publication_year=? WHERE edition_id=?""",
                            (
                                edition["publisher"] or candidate["suggested_publisher"],
                                edition["publication_year"] or candidate["suggested_year"], edition_id,
                            ),
                        )
            connection.execute(
                """UPDATE scan_candidates SET status='approved',proposed_action=?,existing_work_id=?,
                   existing_edition_id=?,existing_file_id=? WHERE candidate_id=?""",
                (action, work_id, edition_id if action not in {"unchanged", "new_version"} else candidate["existing_edition_id"],
                 file_id if action not in {"unchanged", "new_version"} else candidate["existing_file_id"], candidate["candidate_id"]),
            )
            _refresh_search(connection, work_id)
            _sync_work_graph(connection, work_id)
            approved.append({"candidate_id": candidate["candidate_id"], "work_id": work_id, "version_id": version_id})
        remaining = connection.execute(
            "SELECT COUNT(*) FROM scan_candidates WHERE session_id = ? AND status = 'preview'", (session_id,)
        ).fetchone()[0]
        connection.execute(
            "UPDATE scan_sessions SET status = ?, approved_at = ? WHERE session_id = ?",
            ("partially_approved" if remaining else "approved", utc_now(), session_id),
        )
    return {"session_id": session_id, "approved": approved, "remaining": remaining}


def _work_summary(connection: Any, work_id: str) -> dict[str, Any]:
    work = connection.execute("SELECT * FROM works WHERE work_id = ?", (work_id,)).fetchone()
    if work is None:
        raise KeyError(f"unknown work: {work_id}")
    tags = connection.execute(
        """SELECT t.name, wt.origin FROM tags t JOIN work_tags wt ON wt.tag_id = t.tag_id
           WHERE wt.work_id = ? ORDER BY t.name""", (work_id,)
    ).fetchall()
    counts = connection.execute(
        """SELECT COUNT(DISTINCT f.file_id), COUNT(v.version_id) FROM library_files f
           LEFT JOIN file_versions v ON v.file_id = f.file_id WHERE f.work_id = ?""", (work_id,)
    ).fetchone()
    tag_items = [dict(row) for row in tags]
    shelf_tag = next((item["name"] for item in tag_items if item["name"].startswith("shelf:")), "")
    shelf = shelf_tag.removeprefix("shelf:") if shelf_tag else "unclassified"
    return {**dict(work), "tags": tag_items, "shelf": shelf,
            "shelf_label": LIBRARY_SHELVES.get(shelf, shelf),
            "file_count": counts[0], "version_count": counts[1]}


def move_work_to_shelf(project_root: Path, work_id: str, shelf: str,
                       library_root: Path | None = None) -> dict[str, Any]:
    if shelf not in LIBRARY_SHELVES:
        raise ValueError(f"unknown library shelf: {shelf}")
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        if connection.execute("SELECT 1 FROM works WHERE work_id = ?", (work_id,)).fetchone() is None:
            raise KeyError(f"unknown work: {work_id}")
        connection.execute(
            """DELETE FROM work_tags WHERE work_id = ? AND tag_id IN
               (SELECT tag_id FROM tags WHERE name LIKE 'shelf:%')""", (work_id,),
        )
        _add_tag(connection, work_id, f"shelf:{shelf}", "user")
        _refresh_search(connection, work_id)
        _sync_work_graph(connection, work_id)
    return work_detail(project_root, work_id, root)


def search_library(
    project_root: Path,
    query: str = "",
    tags: list[str] | None = None,
    library_root: Path | None = None,
) -> list[dict[str, Any]]:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        if query.strip():
            raw_query = unicodedata.normalize("NFKC", query).strip()
            terms = [term for term in re.split(r"[\s,，;；、/|]+", raw_query) if len(term) > 1]
            terms = list(dict.fromkeys([raw_query, *terms]))
            scores: dict[str, int] = {}
            for index, term in enumerate(terms):
                phrase = '"' + term.replace('"', '""') + '"'
                for row in connection.execute(
                    "SELECT work_id FROM work_search WHERE work_search MATCH ? ORDER BY rank", (phrase,)
                ).fetchall():
                    scores[row["work_id"]] = scores.get(row["work_id"], 0) + (20 if index == 0 else 5)
                contains = f"%{term}%"
                for row in connection.execute(
                    """SELECT DISTINCT w.work_id FROM works w
                       LEFT JOIN editions e ON e.work_id = w.work_id
                       LEFT JOIN library_files f ON f.work_id = w.work_id
                       LEFT JOIN file_versions v ON v.file_id = f.file_id AND v.is_current = 1
                       LEFT JOIN work_tags wt ON wt.work_id = w.work_id
                       LEFT JOIN tags t ON t.tag_id = wt.tag_id
                       WHERE w.canonical_title LIKE ? OR w.author LIKE ? OR e.publisher LIKE ?
                          OR t.name LIKE ? OR v.sample_text LIKE ?""",
                    (contains, contains, contains, contains, contains),
                ).fetchall():
                    scores[row["work_id"]] = scores.get(row["work_id"], 0) + (10 if index == 0 else 2)
            work_ids = sorted(scores, key=lambda work_id: (-scores[work_id], work_id))
        else:
            work_ids = [row["work_id"] for row in connection.execute("SELECT work_id FROM works ORDER BY updated_at DESC")]
        required_tags = set(tags or [])
        results = [_work_summary(connection, work_id) for work_id in work_ids]
    if required_tags:
        results = [item for item in results if required_tags <= {tag["name"] for tag in item["tags"]}]
    return results


def library_assets(
    project_root: Path, kind: str, library_root: Path | None = None,
) -> list[dict[str, Any]]:
    suffixes = {
        "tables": {"xlsx", "xlsm", "csv", "tsv"},
        "maps": {"geojson", "gpkg", "kml", "kmz", "mbtiles"},
        "images": {"png", "jpg", "jpeg", "webp", "gif", "tif", "tiff"},
    }.get(kind)
    if suffixes is None:
        raise ValueError("unknown library asset category")
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        rows = connection.execute(
            """SELECT w.work_id,w.canonical_title,w.author,f.file_id,f.path,v.format,v.byte_count
               FROM works w JOIN library_files f ON f.work_id=w.work_id
               JOIN file_versions v ON v.file_id=f.file_id AND v.is_current=1
               ORDER BY w.updated_at DESC,f.path"""
        ).fetchall()
    return [
        {**dict(row), "available": Path(row["path"]).is_file()}
        for row in rows if str(row["format"]).casefold() in suffixes
    ]


def work_detail(project_root: Path, work_id: str, library_root: Path | None = None) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        result = _work_summary(connection, work_id)
        editions = connection.execute(
            "SELECT * FROM editions WHERE work_id = ? ORDER BY created_at", (work_id,)
        ).fetchall()
        files = connection.execute(
            "SELECT * FROM library_files WHERE work_id = ? ORDER BY path", (work_id,)
        ).fetchall()
        file_items = []
        for file in files:
            path = Path(file["path"])
            current_hash = _file_hash(path) if path.is_file() else ""
            versions = connection.execute(
                "SELECT * FROM file_versions WHERE file_id = ? ORDER BY discovered_at DESC", (file["file_id"],)
            ).fetchall()
            registered_current = next((row for row in versions if row["is_current"]), None)
            file_items.append({
                **dict(file),
                "exists_now": path.is_file(),
                "file_state": (
                    "matches_registered_version"
                    if registered_current and registered_current["sha256"] == current_hash
                    else ("changed_since_last_scan" if path.is_file() else "missing")
                ),
                "versions": [
                    {**dict(row), "bytes_available": bool(row["sha256"] == current_hash)}
                    for row in versions
                ],
            })
        links = connection.execute(
            "SELECT * FROM library_project_links WHERE work_id = ? ORDER BY linked_at", (work_id,)
        ).fetchall()
    with connect(project_root) as connection:
        project_source = connection.execute(
            """SELECT s.source_id,s.title,s.processing_state,s.use_state
               FROM source_library_links l JOIN sources s ON s.source_id=l.source_id
               WHERE l.library_work_id=? ORDER BY l.linked_at DESC LIMIT 1""",
            (work_id,),
        ).fetchone()
    return {
        **result, "editions": [dict(row) for row in editions], "files": file_items,
        "project_links": [dict(row) for row in links],
        "project_source": dict(project_source) if project_source else None,
    }


def update_work(
    project_root: Path,
    work_id: str,
    fields: dict[str, str],
    tags: list[str],
    library_root: Path | None = None,
) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    work_fields = ("canonical_title", "author", "language", "material_type")
    edition_fields = ("edition_label", "publisher", "publication_year", "isbn")
    updates = {key: str(fields[key]).strip() for key in work_fields if key in fields}
    edition_updates = {key: str(fields[key]).strip() for key in edition_fields if key in fields}
    if "canonical_title" in updates and not updates["canonical_title"]:
        raise ValueError("canonical_title cannot be empty")
    with connect_library(root) as connection:
        if connection.execute("SELECT 1 FROM works WHERE work_id = ?", (work_id,)).fetchone() is None:
            raise KeyError(f"unknown work: {work_id}")
        if updates:
            assignments = ", ".join(f"{key} = ?" for key in updates)
            connection.execute(
                f"UPDATE works SET {assignments}, updated_at = ? WHERE work_id = ?",
                [*updates.values(), utc_now(), work_id],
            )
        if edition_updates:
            edition_id = str(fields.get("edition_id", ""))
            if not edition_id:
                edition = connection.execute(
                    "SELECT edition_id FROM editions WHERE work_id = ? ORDER BY created_at LIMIT 1", (work_id,)
                ).fetchone()
                edition_id = edition["edition_id"]
            assignments = ", ".join(f"{key} = ?" for key in edition_updates)
            connection.execute(
                f"UPDATE editions SET {assignments} WHERE edition_id = ? AND work_id = ?",
                [*edition_updates.values(), edition_id, work_id],
            )
        connection.execute(
            """DELETE FROM work_tags WHERE work_id = ? AND origin = 'user'
               AND tag_id NOT IN (SELECT tag_id FROM tags WHERE name LIKE 'shelf:%')""", (work_id,)
        )
        for tag in tags:
            _add_tag(connection, work_id, tag, "user")
        _refresh_search(connection, work_id)
        _sync_work_graph(connection, work_id)
    return work_detail(project_root, work_id, root)


LITERATURE_RELATION_TYPES = {"cites", "uses_material_from", "reviews", "translates", "mentions_work"}


def _derived_literature_relation(block_type: str, text: str) -> str:
    folded = text.casefold()
    if any(term in folded for term in ("翻译", "译本", "译自", "translation of", "translated")):
        return "translates"
    if any(term in folded for term in ("评介", "书评", "评述", "review of", "reviewed")):
        return "reviews"
    if any(term in folded for term in ("材料据自", "材料来自", "采用其材料", "转引自", "uses material from")):
        return "uses_material_from"
    if block_type in {"footnote", "bottom_note", "note", "bibliography", "reference"}:
        return "cites"
    return "mentions_work"


def _literature_relation_candidates(
    project_root: Path, library_connection: Any, query: str = "", limit: int = 300,
) -> list[dict[str, Any]]:
    works = [dict(row) for row in library_connection.execute(
        """SELECT w.work_id, w.canonical_title, w.author,
                  COALESCE((
                      SELECT substr(t.name,7) FROM work_tags wt JOIN tags t ON t.tag_id=wt.tag_id
                      WHERE wt.work_id=w.work_id AND t.name LIKE 'shelf:%' LIMIT 1
                  ),'unclassified') AS shelf
           FROM works w
           WHERE NOT EXISTS (
             SELECT 1 FROM work_tags wt JOIN tags t ON t.tag_id=wt.tag_id
             WHERE wt.work_id=w.work_id AND t.name='shelf:reading_notes'
           ) ORDER BY w.updated_at DESC"""
    ).fetchall() if len(str(row["canonical_title"]).strip()) >= 4
        and row["shelf"] in GRAPH_WORK_SHELVES]
    work_lookup = {item["work_id"]: item for item in works}
    with connect(project_root) as connection:
        blocks = connection.execute(
            """SELECT l.library_work_id AS source_work_id, s.source_id, s.title AS source_title,
                      p.page_id, p.physical_page, p.printed_page, b.block_id, b.block_type,
                      COALESCE(b.human_text, b.machine_text) AS text
               FROM source_library_links l JOIN sources s ON s.source_id = l.source_id
               JOIN pages p ON p.source_id = s.source_id JOIN blocks b ON b.page_id = p.page_id
               WHERE b.use_state = 'research_usable'
               ORDER BY s.created_at, p.physical_page, b.block_order LIMIT 5000"""
        ).fetchall()
        decisions = {
            row["relation_key"]: dict(row) for row in connection.execute(
                "SELECT * FROM literature_relation_decisions"
            ).fetchall()
        }
    candidates: list[dict[str, Any]] = []
    for block in blocks:
        text = " ".join(str(block["text"] or "").split())
        folded = text.casefold()
        if not text:
            continue
        source_work = work_lookup.get(block["source_work_id"], {"canonical_title": block["source_title"], "author": ""})
        for target in works:
            if target["work_id"] == block["source_work_id"]:
                continue
            title = str(target["canonical_title"]).strip()
            if title.casefold() not in folded:
                continue
            relation_key = hashlib.sha256(
                f"{block['source_work_id']}\0{target['work_id']}\0{block['block_id']}".encode("utf-8")
            ).hexdigest()
            decision = decisions.get(relation_key, {})
            derived_type = _derived_literature_relation(block["block_type"], text)
            item = {
                "relation_key": relation_key,
                "source_work_id": block["source_work_id"], "source_work_title": source_work["canonical_title"],
                "target_work_id": target["work_id"], "target_work_title": title,
                "target_author": target["author"], "source_id": block["source_id"],
                "page_id": block["page_id"], "block_id": block["block_id"],
                "physical_page": block["physical_page"], "printed_page": block["printed_page"] or "",
                "quote": text[:1200], "origin": "exact_registered_title_in_markdown",
                "status": decision.get("status", "derived"),
                "relation_type": decision.get("relation_type", derived_type),
                "decided_by": decision.get("decided_by", ""),
                "decision_reason": decision.get("decision_reason", ""),
            }
            if query and query.casefold() not in " ".join((
                item["source_work_title"], item["target_work_title"], item["target_author"], item["quote"],
            )).casefold():
                continue
            candidates.append(item)
            if len(candidates) >= limit:
                return candidates
    return candidates


def decide_literature_relation(
    project_root: Path, relation_key: str, approved: bool, relation_type: str,
    reviewer: str, reason: str, library_root: Path | None = None,
) -> dict[str, Any]:
    relation_key, reviewer, reason = (str(value).strip() for value in (relation_key, reviewer, reason))
    if not relation_key or not reviewer or not reason:
        raise ValueError("relation key, reviewer and reason are required")
    if relation_type not in LITERATURE_RELATION_TYPES:
        raise ValueError(f"unsupported literature relation type: {relation_type}")
    root = library_root_for(project_root, library_root)
    with connect_library(root) as library_connection:
        candidate = next((item for item in _literature_relation_candidates(
            project_root, library_connection, limit=5000
        ) if item["relation_key"] == relation_key), None)
    if candidate is None:
        raise KeyError("literature relation candidate no longer matches the current source text")
    status = "approved" if approved else "rejected"
    now = utc_now()
    with connect(project_root) as connection:
        connection.execute(
            """INSERT INTO literature_relation_decisions(
                   relation_key, source_work_id, target_work_id, relation_type, source_id,
                   page_id, block_id, quote, status, origin, decided_by, decision_reason,
                   created_at, decided_at
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(relation_key) DO UPDATE SET relation_type=excluded.relation_type,
                   quote=excluded.quote, status=excluded.status, origin=excluded.origin,
                   decided_by=excluded.decided_by, decision_reason=excluded.decision_reason,
                   decided_at=excluded.decided_at""",
            (relation_key, candidate["source_work_id"], candidate["target_work_id"], relation_type,
             candidate["source_id"], candidate["page_id"], candidate["block_id"], candidate["quote"],
             status, candidate["origin"], reviewer, reason, now, now),
        )
    return {**candidate, "status": status, "relation_type": relation_type,
            "decided_by": reviewer, "decision_reason": reason}


def library_graph(project_root: Path, query: str = "", limit: int = 200,
                  library_root: Path | None = None, include_reading_notes: bool = False,
                  shelf: str = "") -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    limit = max(1, min(int(limit), 500))
    selected_shelf = shelf if shelf in LIBRARY_SHELVES else ""
    content_graph = project_content_graph(project_root, query, max(40, limit * 5))
    with connect_library(root) as connection:
        literature_relations = _literature_relation_candidates(project_root, connection, query, max(60, limit * 4))
        missing_work_ids = [
            row["work_id"]
            for row in connection.execute(
                """SELECT w.work_id FROM works w
                   LEFT JOIN knowledge_nodes n
                     ON n.node_type = 'work' AND n.normalized_label = lower(w.work_id)
                   WHERE n.node_id IS NULL"""
            ).fetchall()
        ]
        for work_id in missing_work_ids:
            _sync_work_graph(connection, work_id)
        if query.strip():
            aliases = [query.strip()]
            alias_row = connection.execute(
                "SELECT canonical_name FROM author_aliases WHERE alias_normalized=?", (_author_key(query.strip()),)
            ).fetchone()
            if alias_row:
                aliases.extend(row["alias"] for row in connection.execute(
                    "SELECT alias FROM author_aliases WHERE canonical_name=?", (alias_row["canonical_name"],)
                ).fetchall())
            work_ids = []
            for value in dict.fromkeys(aliases):
                phrase = '"' + value.replace('"', '""') + '"'
                work_ids.extend(row["work_id"] for row in connection.execute(
                    "SELECT work_id FROM work_search WHERE work_search MATCH ? LIMIT ?", (phrase, limit)
                ).fetchall() if row["work_id"] not in work_ids)
                contains = f"%{value}%"
                fallback = connection.execute(
                    """SELECT DISTINCT w.work_id FROM works w
                       LEFT JOIN editions e ON e.work_id = w.work_id
                       LEFT JOIN library_files f ON f.work_id = w.work_id
                       LEFT JOIN file_versions v ON v.file_id = f.file_id AND v.is_current = 1
                       LEFT JOIN work_tags wt ON wt.work_id = w.work_id
                       LEFT JOIN tags t ON t.tag_id = wt.tag_id
                       WHERE w.canonical_title LIKE ? OR w.author LIKE ? OR e.publisher LIKE ?
                          OR t.name LIKE ? OR v.sample_text LIKE ? LIMIT ?""",
                    (contains, contains, contains, contains, contains, limit),
                ).fetchall()
                work_ids.extend(row["work_id"] for row in fallback if row["work_id"] not in work_ids)
            work_ids = work_ids[:limit]
        else:
            work_ids = [row["work_id"] for row in connection.execute(
                "SELECT w.work_id FROM works w ORDER BY w.updated_at DESC LIMIT ?", (limit,)
            ).fetchall()]
        if work_ids:
            placeholders = ",".join("?" for _ in work_ids)
            shelves = {row["work_id"]: row["shelf"] for row in connection.execute(
                f"""SELECT w.work_id, COALESCE((
                           SELECT substr(t.name,7) FROM work_tags wt JOIN tags t ON t.tag_id=wt.tag_id
                           WHERE wt.work_id=w.work_id AND t.name LIKE 'shelf:%' LIMIT 1
                       ),'unclassified') AS shelf
                    FROM works w WHERE w.work_id IN ({placeholders})""",
                tuple(work_ids),
            )}
            allowed_shelves = (
                {"reading_notes"} if selected_shelf == "reading_notes"
                else ({selected_shelf} & set(GRAPH_WORK_SHELVES)) if selected_shelf
                else set(GRAPH_WORK_SHELVES)
            )
            if include_reading_notes and not selected_shelf:
                allowed_shelves.add("reading_notes")
            work_ids = [work_id for work_id in work_ids if shelves.get(work_id) in allowed_shelves]
        if not work_ids:
            return {
                "nodes": [], "edges": [], "work_cards": [], "content_graph": content_graph,
                "entity_nodes": [], "entity_edges": [],
                "literature_relations": literature_relations,
                "node_count": 0, "edge_count": 0,
                "backfilled_work_count": len(missing_work_ids), "query": query.strip(),
                "shelf": selected_shelf,
            }
        normalized_ids = [work_id.casefold() for work_id in work_ids]
        placeholders = ",".join("?" for _ in normalized_ids)
        if query.strip():
            order = {work_id.casefold(): index for index, work_id in enumerate(work_ids)}
        else:
            order = {work_id.casefold(): index for index, work_id in enumerate(work_ids)}
        nodes = [dict(row) for row in connection.execute(
            f"""SELECT node_id, node_type, label, normalized_label, origin FROM knowledge_nodes
                WHERE node_type = 'work' AND normalized_label IN ({placeholders})""",
            tuple(normalized_ids),
        )]
        nodes.sort(key=lambda item: order.get(item["normalized_label"], limit))
        node_ids = {row["node_id"] for row in nodes}
        placeholders = ",".join("?" for _ in node_ids)
        edges = [dict(row) for row in connection.execute(
            f"""SELECT edge_id, source_node_id, relation, target_node_id, work_id, origin
                FROM knowledge_edges WHERE source_node_id IN ({placeholders})
                   OR target_node_id IN ({placeholders}) ORDER BY relation""",
            (*node_ids, *node_ids),
        )]
        connected_ids = node_ids | {row["source_node_id"] for row in edges} | {row["target_node_id"] for row in edges}
        placeholders = ",".join("?" for _ in connected_ids)
        nodes = [dict(row) for row in connection.execute(
            f"SELECT node_id, node_type, label, normalized_label, origin FROM knowledge_nodes WHERE node_id IN ({placeholders}) ORDER BY node_type, label",
            tuple(connected_ids),
        )]
        work_lookup = {work_id.casefold(): work_id for work_id in work_ids}
        work_categories = {row["work_id"]: row["shelf"] for row in connection.execute(
            f"""SELECT w.work_id, COALESCE((
                       SELECT substr(t.name,7) FROM work_tags wt JOIN tags t ON t.tag_id=wt.tag_id
                       WHERE wt.work_id=w.work_id AND t.name LIKE 'shelf:%' LIMIT 1
                   ),'unclassified') AS shelf
                FROM works w
                WHERE w.work_id IN ({','.join('?' for _ in work_ids)})""",
            tuple(work_ids),
        )}
        for node in nodes:
            node["work_id"] = work_lookup.get(node["normalized_label"], "") if node["node_type"] == "work" else ""
            node["graph_category"] = work_categories.get(node["work_id"], "unclassified") if node["work_id"] else node["node_type"]
            node.pop("normalized_label", None)
        entity_nodes = [dict(node) for node in nodes if not (
            node["node_type"] == "tag" and str(node["label"]).startswith(("metadata:", "triage:", "material:", "shelf:"))
        )]
        entity_node_ids = {node["node_id"] for node in entity_nodes}
        entity_edges = [dict(edge) for edge in edges if edge["source_node_id"] in entity_node_ids and edge["target_node_id"] in entity_node_ids]
        work_nodes = {node["node_id"]: node for node in nodes if node["node_type"] == "work"}
        metadata_nodes = {node["node_id"]: node for node in nodes if node["node_type"] != "work"}
        related: dict[str, set[str]] = {}
        for edge in edges:
            if edge["source_node_id"] in work_nodes and edge["target_node_id"] in metadata_nodes:
                related.setdefault(edge["target_node_id"], set()).add(edge["source_node_id"])
            elif edge["target_node_id"] in work_nodes and edge["source_node_id"] in metadata_nodes:
                related.setdefault(edge["source_node_id"], set()).add(edge["target_node_id"])
        direct_edges: dict[str, dict[str, str]] = {}
        for metadata_id, work_node_ids in related.items():
            metadata = metadata_nodes[metadata_id]
            node_type = metadata["node_type"]
            if node_type in {"material_type", "tag", "year"}:
                continue
            ordered = sorted(work_node_ids, key=lambda node_id: work_nodes[node_id]["label"].casefold())
            pairs = [(ordered[left], ordered[right]) for left in range(len(ordered)) for right in range(left + 1, len(ordered))]
            for source_id, target_id in pairs:
                if node_type == "person":
                    relation = "same_author"
                elif node_type == "organization":
                    source_category = work_nodes[source_id]["graph_category"]
                    target_category = work_nodes[target_id]["graph_category"]
                    if source_category == target_category == "academic_articles":
                        relation = "same_journal"
                    elif source_category == target_category and source_category in {
                        "primary_sources", "monographs", "reference_works",
                    }:
                        relation = "same_publisher"
                    else:
                        continue
                else:
                    continue
                edge_id = "KGE_" + hashlib.sha256(f"{source_id}\0{relation}\0{target_id}".encode("utf-8")).hexdigest()[:24]
                direct_edges[edge_id] = {
                    "edge_id": edge_id, "source_node_id": source_id, "relation": relation,
                    "target_node_id": target_id, "work_id": "", "origin": "collapsed_bibliographic_metadata",
                }
        work_node_by_work_id = {node["work_id"]: node["node_id"] for node in work_nodes.values()}
        for relation in literature_relations:
            if relation["status"] not in {"approved", "derived"}:
                continue
            source_id = work_node_by_work_id.get(relation["source_work_id"])
            target_id = work_node_by_work_id.get(relation["target_work_id"])
            if not source_id or not target_id:
                continue
            relation_type = relation["relation_type"]
            edge_id = "KGE_" + hashlib.sha256(f"{source_id}\0{relation_type}\0{target_id}".encode("utf-8")).hexdigest()[:24]
            direct_edges[edge_id] = {
                "edge_id": edge_id, "source_node_id": source_id, "relation": relation_type,
                "target_node_id": target_id, "work_id": relation["source_work_id"],
                "origin": "approved_literature_relation" if relation["status"] == "approved" else "markdown_derived_literature_relation",
            }
            entity_edges.append(dict(direct_edges[edge_id]))
        nodes = list(work_nodes.values())
        edges = list(direct_edges.values())
        work_cards = []
        for work_id in work_ids:
            work = _work_summary(connection, work_id)
            edition = connection.execute(
                """SELECT edition_label, publisher, publication_year FROM editions
                   WHERE work_id = ? ORDER BY created_at LIMIT 1""", (work_id,),
            ).fetchone()
            current = connection.execute(
                """SELECT v.sample_text, v.inspected_pages, v.page_count, v.format, v.text_layer,
                          v.triage_state, v.qualification, f.file_id, f.path
                   FROM file_versions v JOIN library_files f ON f.file_id = v.file_id
                   WHERE f.work_id = ? AND v.is_current = 1 ORDER BY v.discovered_at DESC LIMIT 1""",
                (work_id,),
            ).fetchone()
            project_link_count = connection.execute(
                "SELECT COUNT(*) FROM library_project_links WHERE work_id = ?", (work_id,)
            ).fetchone()[0]
            excerpt = " ".join(str(current["sample_text"] if current else "").split())[:900]
            work_cards.append({
                "work_id": work_id, "title": work["canonical_title"], "author": work["author"],
                "material_type": work["material_type"], "shelf": work["shelf"],
                "shelf_label": work["shelf_label"], "edition": dict(edition) if edition else {},
                "content_excerpt": excerpt, "preview_pages": int(current["inspected_pages"] or 0) if current else 0,
                "page_count": current["page_count"] if current else None,
                "format": current["format"] if current else "", "text_layer": current["text_layer"] if current else "",
                "triage_state": current["triage_state"] if current else "",
                "qualification": current["qualification"] if current else "",
                "current_file_id": current["file_id"] if current else "",
                "file_available": bool(current and Path(current["path"]).is_file()),
                "project_link_count": project_link_count,
            })
    project_sources: dict[str, dict[str, Any]] = {}
    if work_ids:
        placeholders = ",".join("?" for _ in work_ids)
        with connect(project_root) as connection:
            for row in connection.execute(
                f"""SELECT l.library_work_id, s.source_id, s.title, s.processing_state, s.use_state
                    FROM source_library_links l JOIN sources s ON s.source_id = l.source_id
                    WHERE l.library_work_id IN ({placeholders})""",
                tuple(work_ids),
            ).fetchall():
                project_sources[row["library_work_id"]] = dict(row)
    for card in work_cards:
        card["project_source"] = project_sources.get(card["work_id"])
    return {
        "nodes": nodes,
        "edges": edges,
        "entity_nodes": entity_nodes,
        "entity_edges": entity_edges,
        "work_cards": work_cards,
        "node_count": len(nodes),
        "edge_count": len(edges),
        "backfilled_work_count": len(missing_work_ids),
        "query": query.strip(),
        "shelf": selected_shelf,
        "preview_boundary": "bounded_intake_sample_not_evidence",
        "content_graph": content_graph,
        "literature_relations": literature_relations,
    }


def library_file_path(project_root: Path, file_id: str, library_root: Path | None = None) -> Path:
    root = library_root_for(project_root, library_root)
    with connect_library(root) as connection:
        row = connection.execute("SELECT path FROM library_files WHERE file_id = ?", (file_id,)).fetchone()
    if row is None:
        raise KeyError(f"unknown library file: {file_id}")
    path = Path(row["path"])
    if not path.is_file():
        raise FileNotFoundError(f"library file is no longer available: {path}")
    return path


def link_work_to_project(
    project_root: Path,
    work_id: str,
    library_root: Path | None = None,
) -> dict[str, Any]:
    root = library_root_for(project_root, library_root)
    with connect(project_root) as project_connection:
        current_project_id = project_id(project_connection)
    with connect_library(root) as connection:
        if connection.execute("SELECT 1 FROM works WHERE work_id = ?", (work_id,)).fetchone() is None:
            raise KeyError(f"unknown work: {work_id}")
        connection.execute(
            """INSERT OR IGNORE INTO library_project_links(work_id, project_id, project_root, linked_at)
               VALUES (?, ?, ?, ?)""",
            (work_id, current_project_id, str(project_root.resolve()), utc_now()),
        )
    return work_detail(project_root, work_id, root)
